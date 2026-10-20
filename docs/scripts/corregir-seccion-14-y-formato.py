#!/usr/bin/env python3
"""
Corrige formato de párrafos insertados, numeración del capítulo XIV y listas.

    python docs/scripts/corregir-seccion-14-y-formato.py <ruta del .docx>

--------------------------------------------------------------------------
1. TEXTOS PEGADOS Y CENTRADOS
--------------------------------------------------------------------------
Los párrafos que los scripts anteriores insertaron con `insert_paragraph_before`
nacieron SIN formato de párrafo: ni `w:spacing`, ni `w:jc`, ni `w:rPr` en sus
runs. Quedaban a interlineado simple, sin justificar y sin separación con el
párrafo siguiente — eso es lo que se ve como «texto pegado».

LA PRIMERA CORRECCIÓN FUE PEOR QUE EL PROBLEMA. Buscó como modelo «el primer
párrafo Normal del documento que sí tuviera espaciado» y eso resultó ser el
párrafo 2: el título «UNIVERSIDAD TECNOLÓGICA DEL PERÚ» de la portada, que está
CENTRADO. Clonarlo dejó centrados y con letra de portada los apartados 14.3,
14.4 y 16.8 a 16.10.

El modelo no se elige por «el primero que encuentre», sino por MAYORÍA MEDIDA:
se agrupan los párrafos del documento por su firma de formato (pPr + rPr del
primer run) y se toma la firma más repetida entre los párrafos de prosa larga.
Resultado de esa medición sobre este documento:

    cuerpo (Normal, 63 de 120)   pPr: spacing after=140 line=300 auto; jc=both
                                 rPr: (ninguno — lo pone el estilo)
    subtítulo (Heading 2, 69/74) pPr: pStyle Ttulo2; spacing before=260 after=140
                                 rPr: b, bCs, sz=25

Se clonan AMBOS, pPr y rPr. Clonar sólo el pPr fue el segundo error de la
primera versión: por eso los subtítulos nuevos salían sin negrita y más pequeños
que los demás.

--------------------------------------------------------------------------
2. SUBPUNTOS SIN CORRELACIÓN (capítulo XIV)
--------------------------------------------------------------------------
«14.2. Endpoints principales» acumulaba cuatro asuntos bajo un título que sólo
anuncia uno. Queda:

    14.1  Organización del Back-End             (sin cambios)
    14.2  Endpoints principales                 (sin cambios)
    14.3  Cálculo fiscal centralizado           NUEVO
    14.4  Documentación del contrato de la API  NUEVO
    14.5  Integración con el Front-End          antes 14.3

El índice es un campo TOC automático: recoge los dos subtítulos nuevos al
pulsar F9.

--------------------------------------------------------------------------
3. LISTAS QUE NO EMPIEZAN EN 1
--------------------------------------------------------------------------
Cuatro listas numeradas distintas comparten la misma instancia de numeración
(`numId=2`), de modo que cada una CONTINÚA la cuenta de la anterior en lugar de
empezar por uno. Las conclusiones salían numeradas del 20 al 29, y las cuatro
etapas de la migración, del 16 al 19 — anunciadas en el texto como «cuatro
etapas».

No se convierte el texto en párrafos con números escritos a mano: se crean
instancias de numeración nuevas que apuntan al MISMO `abstractNum` (misma
sangría, misma viñeta, mismo formato) con `startOverride=1`. Word sigue
tratándolas como listas automáticas.

    Causas identificadas      sigue con numId=2   (ya empieza en 1)
    Consecuencias             instancia nueva     empieza en 1
    Etapas de la migración    instancia nueva     empieza en 1
    Conclusiones              instancia nueva     empieza en 1

--------------------------------------------------------------------------
4. CIFRA AMBIGUA
--------------------------------------------------------------------------
«Con 3 361 productos, la pantalla llegó a mostrar "2 en stock"» → «Con el
catálogo de 3 361 productos DE ENTONCES». El dato es correcto como historia
—el fallo ocurrió con ese catálogo—, pero junto al resto del documento, que
habla de 884, parecía una contradicción.

NO SE REGENERA EL DOCUMENTO: se edita el .docx existente para conservar la
portada de la UTP, los estilos, las fuentes, los encabezados y los colores.
"""
from __future__ import annotations

import copy
import re
import sys
from collections import Counter
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

cambios: list[str] = []
fallos: list[str] = []


# ----------------------------------------------------------------- utilidades
def xml_de(elemento) -> str:
    if elemento is None:
        return ""
    s = etree.tostring(elemento, encoding="unicode")
    return re.sub(r' xmlns:[a-zA-Z0-9]+="[^"]*"', "", s)


def firma(parrafo) -> tuple[str, str]:
    pPr = parrafo._p.find(qn("w:pPr"))
    rPr = parrafo.runs[0]._r.find(qn("w:rPr")) if parrafo.runs else None
    return xml_de(pPr), xml_de(rPr)


def modelo_mayoritario(doc, estilo: str, largo_minimo: int):
    """Párrafo cuyo formato es el más repetido del documento para ese estilo.

    Elegir por mayoría y no por «el primero que aparezca» es lo que evita
    tomar como modelo un título de portada o un pie de figura.
    """
    cuenta: Counter = Counter()
    ejemplo: dict = {}
    for p in doc.paragraphs:
        if p.style.name != estilo or not p.runs:
            continue
        if len(p.text.strip()) < largo_minimo:
            continue
        f = firma(p)
        cuenta[f] += 1
        ejemplo.setdefault(f, p)
    if not cuenta:
        return None, 0, 0
    ganadora, n = cuenta.most_common(1)[0]
    return ejemplo[ganadora], n, sum(cuenta.values())


def sin_formato(parrafo) -> bool:
    pPr = parrafo._p.find(qn("w:pPr"))
    return pPr is None or pPr.find(qn("w:spacing")) is None


def aplicar_modelo(parrafo, modelo) -> None:
    """Clona el pPr del modelo y el rPr de su primer run sobre el párrafo."""
    p = parrafo._p
    viejo = p.find(qn("w:pPr"))
    nuevo = copy.deepcopy(modelo._p.find(qn("w:pPr")))
    if viejo is not None:
        p.replace(viejo, nuevo)
    else:
        p.insert(0, nuevo)

    modelo_rPr = modelo.runs[0]._r.find(qn("w:rPr")) if modelo.runs else None
    for run in parrafo.runs:
        actual = run._r.find(qn("w:rPr"))
        if actual is not None:
            run._r.remove(actual)
        if modelo_rPr is not None:
            run._r.insert(0, copy.deepcopy(modelo_rPr))


def reemplazar(parrafo, viejo: str, nuevo: str) -> bool:
    completo = "".join(r.text for r in parrafo.runs)
    if viejo not in completo:
        return False
    parrafo.runs[0].text = completo.replace(viejo, nuevo)
    for r in parrafo.runs[1:]:
        r.text = ""
    return True


def heading_por_texto(doc, texto: str, nivel: str = "Heading 2"):
    """Sólo encabezados reales.

    Buscar por texto a secas encuentra primero la línea del índice —el campo
    TOC ya trae los títulos con su número de página—, que fue exactamente el
    fallo que coló párrafos enteros dentro del índice en una pasada anterior.
    """
    hallados = [p for p in doc.paragraphs
                if p.style.name == nivel and p.text.strip() == texto]
    return hallados[0] if len(hallados) == 1 else None


def parrafo_que_empieza(doc, inicio: str, desde_estilo: str = "Normal"):
    for p in doc.paragraphs:
        if p.style.name == desde_estilo and p.text.strip().startswith(inicio):
            return p
    return None


# ------------------------------------------------------------------ numeración
def numid_de(parrafo) -> str | None:
    pPr = parrafo._p.find(qn("w:pPr"))
    if pPr is None:
        return None
    numPr = pPr.find(qn("w:numPr"))
    if numPr is None:
        return None
    nid = numPr.find(qn("w:numId"))
    return nid.get(qn("w:val")) if nid is not None else None


def nueva_instancia(doc, numid_origen: str) -> str:
    """Instancia de numeración que reutiliza el mismo abstractNum y reinicia en 1."""
    numbering = doc.part.numbering_part.element
    nums = numbering.findall(qn("w:num"))
    abstract = None
    for n in nums:
        if n.get(qn("w:numId")) == numid_origen:
            abstract = n.find(qn("w:abstractNumId")).get(qn("w:val"))
    if abstract is None:
        raise ValueError(f"no existe numId={numid_origen}")

    siguiente = str(max(int(n.get(qn("w:numId"))) for n in nums) + 1)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), siguiente)
    ani = OxmlElement("w:abstractNumId")
    ani.set(qn("w:val"), abstract)
    num.append(ani)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    inicio = OxmlElement("w:startOverride")
    inicio.set(qn("w:val"), "1")
    override.append(inicio)
    num.append(override)
    numbering.append(num)
    return siguiente


def reiniciar_lista(doc, parrafos, etiqueta: str) -> None:
    if not parrafos:
        fallos.append(f"{etiqueta}: lista no localizada")
        return
    origen = numid_de(parrafos[0])
    if origen is None:
        fallos.append(f"{etiqueta}: el primer elemento no es una lista numerada")
        return
    nuevo = nueva_instancia(doc, origen)
    for p in parrafos:
        p._p.find(qn("w:pPr")).find(qn("w:numPr")).find(qn("w:numId")).set(
            qn("w:val"), nuevo
        )
    cambios.append(
        f"{etiqueta}: {len(parrafos)} elementos reinician en 1 "
        f"(numId {origen} → {nuevo})"
    )


def lista_tras(doc, titulo: str, estilo_titulo: str):
    """Elementos consecutivos de lista que siguen al título indicado."""
    indices = [i for i, p in enumerate(doc.paragraphs)
               if p.style.name == estilo_titulo and p.text.strip().startswith(titulo)]
    if len(indices) != 1:
        return []
    salida = []
    for p in doc.paragraphs[indices[0] + 1:]:
        if numid_de(p) is not None:
            salida.append(p)
        elif salida:
            break
    return salida


# ----------------------------------------------------------------------- main
def main(ruta: Path) -> int:
    doc = Document(str(ruta))

    # --------------------------------------------------------------- 1 formato
    modelos = {}
    for estilo, minimo in (("Normal", 200), ("Heading 2", 10)):
        modelo, n, total = modelo_mayoritario(doc, estilo, minimo)
        if modelo is None:
            fallos.append(f"Sin modelo de formato para «{estilo}»")
            continue
        modelos[estilo] = modelo
        print(f"Modelo «{estilo}»: firma mayoritaria en {n} de {total} párrafos")
        print(f"   ejemplo: {modelo.text.strip()[:60]}")
        print(f"   pPr: {xml_de(modelo._p.find(qn('w:pPr')))}")
        print(f"   rPr: {xml_de(modelo.runs[0]._r.find(qn('w:rPr')))}")

    # Un modelo centrado sería un modelo de portada, no de cuerpo: se aborta
    # antes de propagar el error a medio documento.
    for estilo, modelo in modelos.items():
        pPr = modelo._p.find(qn("w:pPr"))
        jc = pPr.find(qn("w:jc")) if pPr is not None else None
        if jc is not None and jc.get(qn("w:val")) == "center":
            fallos.append(f"ABORTADO: el modelo de «{estilo}» está centrado")
            print("\n".join("  !!  " + f for f in fallos))
            return 1

    arreglados: dict[str, list[int]] = {k: [] for k in modelos}
    for i, p in enumerate(doc.paragraphs):
        if p.style.name not in modelos or not p.text.strip():
            continue
        if not sin_formato(p):
            continue
        aplicar_modelo(p, modelos[p.style.name])
        arreglados[p.style.name].append(i)
    for estilo, lista in arreglados.items():
        if lista:
            cambios.append(
                f"Formato de cuerpo restituido en {len(lista)} párrafo(s) «{estilo}»"
            )

    # ------------------------------------------------------------ 2 capítulo XIV
    integracion = heading_por_texto(doc, "14.3. Integración con el Front-End")
    p_docs = parrafo_que_empieza(doc, "El contrato completo de la API")
    p_fiscal = parrafo_que_empieza(doc, "El cálculo del impuesto vive en un único servicio")

    if None in (integracion, p_docs, p_fiscal):
        fallos.append("No se localizaron los anclajes del capítulo XIV")
    else:
        if reemplazar(integracion, "14.3. Integración", "14.5. Integración"):
            cambios.append("Renumerado «14.3. Integración con el Front-End» → 14.5")
        else:
            fallos.append("No se pudo renumerar el apartado de integración")

        for ancla, titulo in (
            (p_docs, "14.4. Documentación del contrato de la API"),
            (p_fiscal, "14.3. Cálculo fiscal centralizado"),
        ):
            nuevo = ancla.insert_paragraph_before(titulo)
            nuevo.style = doc.styles["Heading 2"]
            if "Heading 2" in modelos:
                aplicar_modelo(nuevo, modelos["Heading 2"])
            cambios.append(f"Subtítulo nuevo: «{titulo}»")

    # ---------------------------------------------------------------- 3 listas
    reiniciar_lista(doc, lista_tras(doc, "Consecuencias", "Heading 3"),
                    "Consecuencias del sistema actual")
    reiniciar_lista(doc, lista_tras(doc, "7.3. Migración desde el sistema actual", "Heading 2"),
                    "Etapas de la migración")
    reiniciar_lista(doc, lista_tras(doc, "XX. CONCLUSIONES", "Heading 1"),
                    "Conclusiones")

    # ----------------------------------------------------------------- 4 cifra
    p = parrafo_que_empieza(doc, "Del lado del Front-End se aplicó una regla")
    if p is None:
        fallos.append("No se localizó el párrafo del recuento del catálogo")
    elif reemplazar(p, "Con 3 361 productos, la pantalla",
                    "Con el catálogo de 3 361 productos de entonces, la pantalla"):
        cambios.append("Precisada la cifra histórica del catálogo («de entonces»)")
    else:
        fallos.append("No se encontró «Con 3 361 productos, la pantalla»")

    doc.save(str(ruta))

    print("\nCAMBIOS APLICADOS")
    for c in cambios:
        print("  OK  ", c)
    if fallos:
        print("\nNO APLICADOS (revisar a mano)")
        for f in fallos:
            print("  !!  ", f)
    print(f"\nTotal: {len(cambios)} cambios, {len(fallos)} sin aplicar.")
    return 1 if fallos else 0


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(2)
    sys.exit(main(Path(sys.argv[1])))
