#!/usr/bin/env python3
"""
Lleva TODAS las cifras de verificacion del informe a un unico corte.

    python docs/scripts/unificar-corte-de-cifras.py <ruta del .docx>

EL PROBLEMA QUE RESUELVE

El documento arrastraba cuatro cortes distintos conviviendo:

    parrafo 682 + Figura 43   98 pruebas / 246 aserciones  («paso de 89 a 98»)
    parrafos 624 y 671        165 pruebas / 497 aserciones
    Anexo C.3                 139 pruebas / 421 aserciones  (su imagen decia 98)
    ejecucion real de hoy     198 pruebas / 603 aserciones

Y lo mismo con la base: el parrafo 559 hablaba de 46 migraciones, 32 tablas y
3 361 productos, mientras la Figura 50 y la C.1 decian 45 migraciones y la
Figura C.2 mostraba 884 productos.

Actualizar solo dos figuras habria dejado el documento MENOS coherente, no
mas: el pie de la Figura 43 contradiria al parrafo que tiene justo encima.
Por eso esto toca a la vez las imagenes, los pies y la prosa que los cita.

CORTE QUE MANDA (30/10/2026, PHP 8.4 y PostgreSQL 16)

    php artisan test          198 passed (603 assertions)   13.31 s
    php artisan migrate:status 49 migraciones, todas Ran
    base de desarrollo        34 tablas, 884 productos, 886 lotes
    frontend (vitest)         62 pruebas en 8 archivos

QUE NO SE TOCA

La Figura C.4 (las cinco suites de reglas de negocio) se reejecuto y devuelve
exactamente la misma salida, con exit=0 en las cinco. Su imagen sigue siendo
fiel y se deja como esta.
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from PIL import Image

# --- sustituciones de texto ------------------------------------------------
# (viejo, nuevo, etiqueta). Si alguna no aparece, el script lo grita: significa
# que el documento cambio y esta receta se quedo vieja.
TEXTOS = [
    # Prosa
    ("las mismas 46 migraciones, las mismas 32 tablas y los mismos 3 361 productos con 3 363 lotes",
     "las mismas 49 migraciones, las mismas 34 tablas y los mismos 884 productos con 886 lotes",
     "§15 · base de desarrollo"),

    ("El segundo son 165 pruebas automatizadas",
     "El segundo son 198 pruebas automatizadas",
     "§16.1 · niveles de prueba"),

    ("La suite reúne 165 pruebas con 497 aserciones y se ejecuta completa en poco más de catorce segundos",
     "La suite reúne 198 pruebas con 603 aserciones y se ejecuta completa en poco más de trece segundos",
     "§16.7 · recuento de la suite"),

    # El parrafo 682 narraba un salto puntual (89 -> 98) que ya no describe al
    # documento. Se reescribe para que hable del corte unico, no de un tramo.
    ("La suite pasó de 89 a 98 casos durante esta entrega, y corre sobre PostgreSQL, el mismo motor que producción.",
     "La suite reúne hoy 198 casos con 603 aserciones y corre sobre PostgreSQL, el mismo motor que producción.",
     "§16.10 · crecimiento de la suite"),

    # Tablas
    ("165 pruebas con 497 aserciones, ejecutadas en la integración continua",
     "198 pruebas con 603 aserciones, ejecutadas en la integración continua",
     "Tabla 30 · integración continua"),

    ("165 pruebas y 497 aserciones sobre PostgreSQL, todas en verde.",
     "198 pruebas y 603 aserciones sobre PostgreSQL, todas en verde.",
     "Tabla 62 · E-05"),

    ("165 pruebas, 497 aserciones, 14.05 s",
     "198 pruebas, 603 aserciones, 13.31 s",
     "Tabla 99 · medición de la suite"),

    # Pies de figura
    ("Figura 43. Suite automatizada tras cerrar el acceso de gestión: 98 pruebas y 246 aserciones en verde.",
     "Figura 43. Suite automatizada completa sobre PostgreSQL 16: 198 pruebas y 603 aserciones en verde.",
     "Pie · Figura 43"),

    ("las 45 figuran como ejecutadas y ninguna como pendiente",
     "las 49 figuran como ejecutadas y ninguna como pendiente",
     "Pie · Figura 50"),

    ("las 45 aparecen como aplicadas y ninguna como pendiente",
     "las 49 aparecen como aplicadas y ninguna como pendiente",
     "Pie · Figura C.1"),

    ("Figura C.3. Las 139 pruebas con 421 aserciones ejecutadas sobre PostgreSQL",
     "Figura C.3. Las 198 pruebas con 603 aserciones ejecutadas sobre PostgreSQL",
     "Pie · Figura C.3"),
]

# --- reemplazo de imagenes -------------------------------------------------
# rId -> archivo nuevo
IMAGENES = {
    "rId49": "evidencias/cortes/suite-198.png",       # Figura 43
    "rId115": "evidencias/cortes/suite-198.png",      # Figura C.3
    "rId54": "evidencias/cortes/migraciones-49.png",  # Figura C.1
}

# Pie de figura que se quedo SIN imagen y hay que dotar de una.
PIE_HUERFANO = "Figura 50. Migraciones aplicadas sobre la base de producción"
IMAGEN_HUERFANA = "evidencias/cortes/migraciones-49.png"

cambios: list[str] = []
fallos: list[str] = []


def recorrer(doc):
    for p in doc.paragraphs:
        yield p
    for t in doc.tables:
        for fila in t.rows:
            for celda in fila.cells:
                for p in celda.paragraphs:
                    yield p


def reemplazar(parrafo, viejo: str, nuevo: str) -> bool:
    """Sustituye conservando el formato del primer run.

    Word parte un parrafo en runs por motivos suyos —una correccion
    ortografica, un cambio de revision—, asi que el texto buscado casi nunca
    vive en un solo run. Se recompone, se sustituye y se devuelve todo al
    primer run, que es el que lleva el formato bueno.
    """
    completo = "".join(r.text for r in parrafo.runs)
    if viejo not in completo:
        return False
    parrafo.runs[0].text = completo.replace(viejo, nuevo)
    for r in parrafo.runs[1:]:
        r.text = ""
    return True


def _ajustar_alto(doc, rid: str, proporcion: float) -> int:
    """Recalcula el alto declarado de cada dibujo que use este rId.

    Si solo se cambian los bytes, Word conserva el alto del XML y la imagen
    nueva sale aplastada o estirada. Se mantiene el ancho —que es lo que
    encaja en el margen— y se recalcula el alto con la proporcion real.
    """
    tocados = 0
    for blip in doc.element.body.iter(qn("a:blip")):
        if blip.get(qn("r:embed")) != rid:
            continue
        # El <a:blip> cuelga del <pic:pic>; el tamanio vive en dos sitios que
        # tienen que concordar: el <wp:extent> del contenedor y el <a:ext> de
        # la transformacion. Si solo se cambia uno, Word recorta.
        # Se sube hasta el <w:drawing>: el tamanio vive en DOS sitios, el
        # <wp:extent> del marco y el <a:ext> del dibujo, y tienen que decir lo
        # mismo. Partiendo de un ancestro mas cercano solo se alcanza el
        # segundo, y Word dibuja la imagen a un tamanio y la recorta al otro.
        marco = blip
        while marco is not None and marco.tag != qn("w:drawing"):
            marco = marco.getparent()
        if marco is None:
            continue
        for ext in marco.iter():
            if ext.tag in (qn("wp:extent"), qn("a:ext")):
                cx = ext.get("cx")
                if cx:
                    ext.set("cy", str(int(int(cx) * proporcion)))
                    tocados += 1
    return tocados


def sustituir_imagen(doc, rid: str, archivo: Path) -> bool:
    """Cambia los bytes de una imagen y ajusta su alto."""
    try:
        parte = doc.part.related_parts[rid]
    except KeyError:
        return False

    parte._blob = archivo.read_bytes()
    with Image.open(archivo) as im:
        ancho, alto = im.size
    _ajustar_alto(doc, rid, alto / ancho)
    return True


def dotar_de_imagen(doc, inicio_pie: str, archivo: Path) -> bool:
    """Inserta una imagen delante de un pie de figura que se quedo sin ella.

    La Figura 50 tenia pie pero no figura: alguien borro la imagen y el pie
    sobrevivio. Un pie sin imagen es peor que ninguno de los dos, porque el
    indice de figuras la cuenta y el lector la busca.
    """
    for i, p in enumerate(doc.paragraphs):
        if not p.text.strip().startswith(inicio_pie):
            continue
        hueco = doc.paragraphs[i].insert_paragraph_before()
        hueco.alignment = p.alignment
        ancho = doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin
        hueco.add_run().add_picture(str(archivo), width=ancho)
        return True
    return False


def main(ruta: Path) -> int:
    raiz = Path(__file__).resolve().parents[1]  # docs/
    doc = Document(str(ruta))

    for viejo, nuevo, etiqueta in TEXTOS:
        n = sum(1 for p in recorrer(doc) if reemplazar(p, viejo, nuevo))
        (cambios if n else fallos).append(
            f"{etiqueta}: {n} sustitución(es)" if n
            else f"{etiqueta}: NO SE ENCONTRÓ «{viejo[:55]}…»"
        )

    for rid, relativo in IMAGENES.items():
        archivo = raiz / relativo
        if sustituir_imagen(doc, rid, archivo):
            cambios.append(f"Imagen {rid} ← {relativo}")
        else:
            fallos.append(f"Imagen {rid}: NO EXISTE esa relación en el documento")

    if dotar_de_imagen(doc, PIE_HUERFANO, raiz / IMAGEN_HUERFANA):
        cambios.append(f"Figura 50: imagen insertada ← {IMAGEN_HUERFANA}")
    else:
        fallos.append("Figura 50: no se encontró el pie huérfano")

    doc.save(str(ruta))

    print("CAMBIOS")
    for c in cambios:
        print("  OK  ", c)
    if fallos:
        print("\nNO APLICADO")
        for f in fallos:
            print("  !!  ", f)
    print(f"\nTotal: {len(cambios)} aplicados, {len(fallos)} sin aplicar.")
    return 1 if fallos else 0


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(2)
    sys.exit(main(Path(sys.argv[1])))
