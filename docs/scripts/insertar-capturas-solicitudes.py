# -*- coding: utf-8 -*-
"""
Inserta en el apartado 21.3 del informe la captura de la issue de cada
solicitud de servicio, y renumera TODAS las figuras por orden de aparición.

QUE HACE

1. Para cada subsección SOL-001 a SOL-006 añade, al final de su desarrollo,
   la captura de la issue correspondiente de GitHub con su pie de figura y la
   fuente. La imagen se inserta con `add_picture`, que crea una parte de
   imagen propia por archivo; clonar un párrafo de figura existente no habría
   servido, porque el clon reutiliza el mismo identificador de relación y
   mostraría siempre la misma imagen.

2. Renumera las figuras siguiendo el orden en que aparecen en el documento y
   reescribe las referencias cruzadas del texto con el mismo mapa. Esto
   corrige, además de absorber las seis figuras nuevas, dos desórdenes
   heredados del avance anterior: las figuras 48, 49 y 50 aparecían después
   de la 52 y la 53, y la 51 estaba en el anexo, al final del documento.

POR QUE LA RENUMERACION ES SEGURA

Se construye el mapa antes de tocar nada y las sustituciones de referencias
se hacen en una sola pasada con una función de reemplazo, no encadenando
reemplazos: así «Figura 48 -> 50» y «Figura 50 -> 52» no se pisan entre sí.
Al final se comprueba que los pies van del 1 al N sin huecos ni repeticiones
y que ninguna referencia apunta a un número inexistente.
"""

import copy
import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.table import Table
from docx.text.paragraph import Paragraph

if len(sys.argv) < 3:
    raise SystemExit('Uso: python insertar-capturas-solicitudes.py <docx> <carpeta de capturas>')

RUTA = Path(sys.argv[1])
CAPTURAS = Path(sys.argv[2])
ANCHO = Cm(13)

FIGURAS = [
    ('SOL-001', 'E20-issue-sol001.jpg',
     'Issue #14 en GitHub: ficha de la solicitud SOL-001 con sus campos de seguimiento, '
     'las etiquetas aplicadas y el estado cerrado.'),
    ('SOL-002', 'E21-issue-sol002.jpg',
     'Issue #15 en GitHub: ficha de la solicitud SOL-002, con la explicación entre stock '
     'contable y stock vendible como respuesta registrada.'),
    ('SOL-003', 'E22-issue-sol003.jpg',
     'Issue #16 en GitHub: ficha de la solicitud SOL-003, con la aprobación del propietario '
     'y el rol vendedor asignado por mínimo privilegio.'),
    ('SOL-004', 'E23-issue-sol004.jpg',
     'Issue #17 en GitHub: ficha de la solicitud SOL-004, denegada por mínimo privilegio y '
     'cerrada como «not planned» con el motivo documentado.'),
    ('SOL-005', 'E24-issue-sol005.jpg',
     'Issue #18 en GitHub: ficha de la solicitud SOL-005, atendida por el modelo estándar '
     '«Restablecer contraseña» sin registrar en ningún sitio el valor de la contraseña.'),
    ('SOL-006', 'E25-issue-sol006.jpg',
     'Issue #19 en GitHub: ficha de la solicitud SOL-006, con el respaldo ejecutado antes de '
     'la carga y verificado por conteo de filas.'),
]


def bloques(doc):
    salida = []
    for hijo in doc.element.body.iterchildren():
        if hijo.tag == qn('w:p'):
            salida.append(Paragraph(hijo, doc))
        elif hijo.tag == qn('w:tbl'):
            salida.append(Table(hijo, doc))
    return salida


def poner_texto(parrafo, texto):
    runs = parrafo.runs
    if not runs:
        parrafo.add_run(texto)
        return
    runs[0].text = texto
    for sobrante in runs[1:]:
        sobrante._element.getparent().remove(sobrante._element)


ES_PIE = re.compile(r'^Figura (\d+)\.')


def pies(doc):
    return [p for p in bloques(doc)
            if isinstance(p, Paragraph) and ES_PIE.match(p.text.strip())]


def main():
    doc = Document(str(RUTA))
    items = bloques(doc)

    # modelos: el párrafo que contiene una imagen y su pie (los de la Figura 60)
    i_pie = next(i for i, it in enumerate(items)
                 if isinstance(it, Paragraph) and it.text.strip().startswith('Figura 60.'))
    modelo_pie = items[i_pie]
    modelo_imagen = items[i_pie - 1]

    # --- 1. inserción de las seis capturas --------------------------------
    insertadas = 0
    for codigo, archivo, leyenda in FIGURAS:
        ruta_imagen = CAPTURAS / archivo
        if not ruta_imagen.exists():
            raise SystemExit(f'No se encuentra la captura {ruta_imagen}')

        items = bloques(doc)
        i_sub = next(i for i, it in enumerate(items)
                     if isinstance(it, Paragraph)
                     and it.style.name.startswith('Heading 3')
                     and it.text.strip().startswith(codigo))
        # final de la subsección: el siguiente encabezado o el siguiente capítulo
        i_fin = next(i for i in range(i_sub + 1, len(items))
                     if isinstance(items[i], Paragraph)
                     and items[i].style.name.startswith(('Heading 1', 'Heading 2', 'Heading 3')))
        destino = items[i_fin]._p

        # la imagen se crea al final del documento y luego se mueve
        parrafo_img = doc.add_paragraph()
        parrafo_img._p.addprevious(copy.deepcopy(modelo_imagen._p))  # marcador de formato
        marcador = parrafo_img._p.getprevious()
        parrafo_img.add_run().add_picture(str(ruta_imagen), width=ANCHO)

        # se adopta el formato del párrafo modelo (centrado, espaciado)
        pPr_modelo = marcador.find(qn('w:pPr'))
        if pPr_modelo is not None:
            viejo = parrafo_img._p.find(qn('w:pPr'))
            if viejo is not None:
                parrafo_img._p.remove(viejo)
            parrafo_img._p.insert(0, copy.deepcopy(pPr_modelo))
        marcador.getparent().remove(marcador)

        nuevo_pie = copy.deepcopy(modelo_pie._p)
        destino.addprevious(parrafo_img._p)
        destino.addprevious(nuevo_pie)
        poner_texto(Paragraph(nuevo_pie, doc),
                    f'Figura {9000 + insertadas}. {leyenda} Fuente: Elaboración propia.')
        insertadas += 1

    print(f'Capturas insertadas: {insertadas}')

    # --- 2. mapa de renumeración ------------------------------------------
    lista = pies(doc)
    mapa = {}
    for posicion, p in enumerate(lista, start=1):
        antes = int(ES_PIE.match(p.text.strip()).group(1))
        if antes < 9000:            # las nuevas llevan 9000..9005 y no se mapean
            mapa[antes] = posicion

    movidas = {a: n for a, n in mapa.items() if a != n}
    print(f'Figuras que cambian de número: {len(movidas)}')
    for a, n in sorted(movidas.items()):
        print(f'   Figura {a} -> Figura {n}')

    # --- 3. reescritura de los pies ---------------------------------------
    for posicion, p in enumerate(lista, start=1):
        texto = p.text.strip()
        poner_texto(p, ES_PIE.sub(f'Figura {posicion}.', texto, count=1))

    # --- 4. reescritura de las referencias cruzadas ------------------------
    def recorrer(doc):
        for p in doc.paragraphs:
            yield p
        for t in doc.tables:
            for fila in t.rows:
                for celda in fila.cells:
                    for p in celda.paragraphs:
                        yield p

    cambios_ref = 0
    for p in recorrer(doc):
        texto = p.text
        if 'Figura' not in texto or ES_PIE.match(texto.strip()):
            continue

        def sustituir(m):
            n = int(m.group(1))
            return f'Figura {mapa.get(n, n)}'

        nuevo = re.sub(r'Figura (\d+)', sustituir, texto)
        if nuevo != texto:
            poner_texto(p, nuevo)
            cambios_ref += 1
    print(f'Párrafos con referencias cruzadas actualizadas: {cambios_ref}')

    # --- 5. comprobación ---------------------------------------------------
    numeros = [int(ES_PIE.match(p.text.strip()).group(1)) for p in pies(doc)]
    if numeros != list(range(1, len(numeros) + 1)):
        raise SystemExit(f'Los pies no quedaron correlativos: {numeros}')
    validos = set(numeros)
    huerfanas = set()
    for p in recorrer(doc):
        t = p.text.strip()
        if ES_PIE.match(t):
            continue
        for m in re.finditer(r'Figura (\d+)', t):
            if int(m.group(1)) not in validos:
                huerfanas.add(int(m.group(1)))
    if huerfanas:
        raise SystemExit(f'Referencias a figuras inexistentes: {sorted(huerfanas)}')

    doc.save(str(RUTA))
    print(f'Guardado. Figuras totales: {len(numeros)} (1 a {len(numeros)}), en orden.')


if __name__ == '__main__':
    main()
