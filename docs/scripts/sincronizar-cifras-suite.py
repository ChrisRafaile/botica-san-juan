#!/usr/bin/env python3
"""
Pone las cifras de la suite de pruebas del .docx al dia.

    python docs/scripts/sincronizar-cifras-suite.py <ruta del .docx>

MEDICION QUE MANDA (21/10/2026, sobre PHP 8.4.27 y PostgreSQL 16)

    php artisan test
    Tests:    150 passed (458 assertions)
    Duration: 17.79s

El informe arrastraba 139 pruebas con 421 aserciones, que fue cierto hasta que
la suite creció con los casos de trazabilidad y de checkout. Tambien decia que
corria «en menos de doce segundos», cifra que ya no se sostiene: la ejecucion
medida tarda 17.79 s.

QUE SE CAMBIA Y QUE NO

Se actualiza toda afirmacion sobre el ESTADO ACTUAL de la suite. NO se tocan
los pies de figura:

    Figura 43   «98 pruebas y 246 aserciones»     se queda
    Figura C.3  «139 pruebas con 421 aserciones»  se queda

Un pie describe lo que se ve en una captura tomada en un momento concreto.
Reescribirlo para que cuadre con la cifra de hoy haria que el pie mintiera
sobre su propia imagen: el lector abre la figura, cuenta y ve otra cosa. Esas
dos figuras habria que RECAPTURARLAS, que es un trabajo distinto de editar
texto, y queda anotado como pendiente en lugar de disimularse.
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document

SUSTITUCIONES = [
    ("El segundo son 139 pruebas automatizadas",
     "El segundo son 150 pruebas automatizadas",
     "16.1 · alcance de las pruebas"),

    ("La suite reúne 139 pruebas con 421 aserciones y se ejecuta completa en menos "
     "de doce segundos",
     "La suite reúne 150 pruebas con 458 aserciones y se ejecuta completa en menos "
     "de veinte segundos",
     "16.7 · recuento y duración de la suite"),

    ("139 pruebas con 421 aserciones, ejecutadas en la integración continua",
     "150 pruebas con 458 aserciones, ejecutadas en la integración continua",
     "Tabla 30 · pruebas en integración continua"),

    ("139 pruebas y 421 aserciones sobre PostgreSQL, todas en verde.",
     "150 pruebas y 458 aserciones sobre PostgreSQL, todas en verde.",
     "Tabla 62 · E-05, suite automatizada"),
]

cambios: list[str] = []
fallos: list[str] = []


def reemplazar(parrafo, viejo: str, nuevo: str) -> bool:
    completo = "".join(r.text for r in parrafo.runs)
    if viejo not in completo:
        return False
    parrafo.runs[0].text = completo.replace(viejo, nuevo)
    for r in parrafo.runs[1:]:
        r.text = ""
    return True


def recorrer(doc):
    for p in doc.paragraphs:
        yield p
    for t in doc.tables:
        for fila in t.rows:
            for celda in fila.cells:
                for p in celda.paragraphs:
                    yield p


def main(ruta: Path) -> int:
    doc = Document(str(ruta))

    for viejo, nuevo, etiqueta in SUSTITUCIONES:
        n = sum(1 for p in recorrer(doc) if reemplazar(p, viejo, nuevo))
        if n:
            cambios.append(f"{etiqueta}: {n} sustitución(es)")
        else:
            fallos.append(f"{etiqueta}: NO SE ENCONTRÓ «{viejo[:60]}»")

    doc.save(str(ruta))

    print("CAMBIOS APLICADOS")
    for c in cambios:
        print("  OK  ", c)
    if fallos:
        print("\nNO APLICADOS")
        for f in fallos:
            print("  !!  ", f)
    print("\nPENDIENTE: las Figuras 43 y C.3 muestran ejecuciones anteriores "
          "(98/246 y 139/421). Hay que recapturarlas para que coincidan.")
    print(f"Total: {len(cambios)} cambios, {len(fallos)} sin aplicar.")
    return 1 if fallos else 0


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(2)
    sys.exit(main(Path(sys.argv[1])))
