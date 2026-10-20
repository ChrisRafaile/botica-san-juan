#!/usr/bin/env python3
"""
Sincroniza con la realidad las cifras del .docx que quedaron desfasadas.

    python docs/scripts/sincronizar-cifras-informe.py <ruta del .docx>

CÓMO APARECIERON

El informe se ha ido actualizando por capítulos. Cada vez que creció el
proyecto se corrigieron las cifras del capítulo que se estaba tocando, pero las
mismas cifras repetidas en tablas de otros capítulos se quedaron atrás. El
renderizado completo a PDF las puso a la vista: el mismo documento afirmaba 31,
41 y 45 migraciones en páginas distintas.

MEDICIONES CONTRA EL REPOSITORIO (19/10/2026)

    migraciones   ls botica-san-juan-backend/database/migrations/*.php   45
    commits       git rev-list --count origin/main                      113
    suite         declarada en 16.7 y en la Figura C.3      139 pruebas / 421

QUÉ SE CORRIGE Y QUÉ NO

Se corrige toda afirmación sobre el ESTADO ACTUAL del proyecto. No se toca
ninguna afirmación HISTÓRICA, porque no son errores sino el relato de lo que
pasó:

    «La arquitectura declaraba 16 controladores, 31 migraciones...»   se queda
    «El informe anterior afirmaba 89 pruebas mientras cinco...»       se queda
    pie de la Figura 43: «98 pruebas y 246 aserciones»                se queda
    pie de la Figura 50: «43 commits y 134 archivos»                  se queda

Los pies de figura describen lo que se ve en una captura tomada en un momento
concreto. Cambiar el pie para que coincida con la cifra de hoy haría que el pie
mintiera sobre su propia imagen. Lo que se corrige es al revés: las celdas que
citan esas figuras con cifras distintas a las que la figura muestra.

QUEDA PENDIENTE Y NO SE PUEDE ARREGLAR DESDE AQUÍ

La Figura 11 (modelo de datos conceptual) lleva rotulado «26 migraciones
versionadas» DENTRO de la imagen. Es un PNG: habría que regenerar el diagrama.
Se deja anotado en lugar de disimularlo.
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document

# (texto viejo, texto nuevo, por qué)
SUSTITUCIONES = [
    ("se aplicaron las 31 migraciones",
     "se aplicaron las 45 migraciones",
     "Tabla 23 · migraciones aplicadas en desarrollo"),

    ("Actualizadas a 20 controladores, 41 migraciones y 22 modelos.",
     "Actualizadas a 20 controladores, 45 migraciones y 22 modelos.",
     "Tabla 41 · cierre de la observación O-03"),

    ("41 migraciones en el repositorio;",
     "45 migraciones en el repositorio;",
     "Tabla 47 · versionado del esquema"),

    ("Historial de 45 commits atómicos en español,",
     "Historial de más de 100 commits atómicos en español,",
     "Tabla 62 · E-01, historial del repositorio"),

    ("Revisión del incremento de los Sprints 3 y 4: 41 commits, 131 archivos,",
     "Revisión del incremento de los Sprints 3 y 4: 43 commits, 134 archivos,",
     "Tabla 62 · E-02, para que coincida con la Figura 50 que cita"),

    ("41 migraciones aplicadas sobre PostgreSQL 16 gestionado,",
     "45 migraciones aplicadas sobre PostgreSQL 16 gestionado,",
     "Tabla 62 · E-03, migraciones en producción"),

    ("98 pruebas y 246 aserciones sobre PostgreSQL, todas en verde.",
     "139 pruebas y 421 aserciones sobre PostgreSQL, todas en verde.",
     "Tabla 62 · E-05, suite automatizada actual"),

    ("El segundo son 98 pruebas automatizadas que protegen contra regresiones.",
     "El segundo son 139 pruebas automatizadas que protegen contra regresiones.",
     "16.1 · alcance de las pruebas"),
]

cambios: list[str] = []
fallos: list[str] = []


def reemplazar(parrafo, viejo: str, nuevo: str) -> bool:
    completo = "".join(r.text for r in parrafo.runs)
    if viejo not in completo:
        return False
    parrafo.runs[0].text = completo.replace(viejo, nuevo)
    for r in parrafo.runs[1:]:
        r.text = ""
    return True


def recorrer(doc):
    for p in doc.paragraphs:
        yield p
    for t in doc.tables:
        for fila in t.rows:
            for celda in fila.cells:
                for p in celda.paragraphs:
                    yield p


def main(ruta: Path) -> int:
    doc = Document(str(ruta))

    for viejo, nuevo, etiqueta in SUSTITUCIONES:
        n = sum(1 for p in recorrer(doc) if reemplazar(p, viejo, nuevo))
        if n:
            cambios.append(f"{etiqueta}: {n} sustitución(es)")
        else:
            fallos.append(f"{etiqueta}: NO SE ENCONTRÓ «{viejo[:60]}»")

    doc.save(str(ruta))

    print("CAMBIOS APLICADOS")
    for c in cambios:
        print("  OK  ", c)
    if fallos:
        print("\nNO APLICADOS")
        for f in fallos:
            print("  !!  ", f)
    print("\nPENDIENTE: la Figura 11 rotula «26 migraciones versionadas» dentro "
          "del PNG; hay que regenerar el diagrama para corregirlo.")
    print(f"Total: {len(cambios)} cambios, {len(fallos)} sin aplicar.")
    return 1 if fallos else 0


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(2)
    sys.exit(main(Path(sys.argv[1])))
