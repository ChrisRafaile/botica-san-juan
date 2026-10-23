#!/usr/bin/env python3
"""
Incorpora los capítulos del APF3 al .docx existente.

    python docs/scripts/volcar-apf3-al-docx.py <docx> <carpeta de capturas>

POR QUÉ SE EDITA EL DOCUMENTO Y NO SE GENERA UNO NUEVO

La portada de la UTP, los estilos institucionales, los encabezados, los
colores corporativos y el campo TOC viven en ese archivo. Regenerarlo con
Pandoc o con python-docx desde cero produciría un documento que dice lo mismo
y no se parece en nada, y habría que rehacer a mano todo el formato.

DÓNDE SE INSERTA Y POR QUÉ AHÍ

El documento termina con XX. CONCLUSIONES, XXI. REFERENCIAS y los anexos. Los
capítulos nuevos van ANTES de las conclusiones, no después: unas conclusiones
seguidas de seis capítulos más dejarían de ser conclusiones. Eso obliga a
renumerar los dos capítulos finales, que pasan a XXVI y XXVII.

El índice es un campo TOC automático, así que recoge los títulos nuevos y la
renumeración al pulsar F9. No hay que tocarlo.

FORMATO

Cada párrafo y cada título se crean clonando el `w:pPr` y el `w:rPr` del
párrafo mayoritario de su estilo, por la misma razón que en el resto de
parches: un párrafo insertado sin formato nace a interlineado simple y sin
separación, y se lee como texto pegado al siguiente.
"""
from __future__ import annotations

import copy
import re
import sys
from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches
from lxml import etree
from docx.oxml.ns import qn

cambios: list[str] = []
fallos: list[str] = []


# ------------------------------------------------------------------ formato
def xml_de(elemento) -> str:
    if elemento is None:
        return ""
    return re.sub(r' xmlns:[a-zA-Z0-9]+="[^"]*"', "", etree.tostring(elemento, encoding="unicode"))


def modelo_de(doc, estilo: str, largo_minimo: int):
    """Párrafo cuyo formato es el más repetido del documento para ese estilo."""
    cuenta: Counter = Counter()
    ejemplo: dict = {}

    for p in doc.paragraphs:
        if p.style.name != estilo or not p.runs or len(p.text.strip()) < largo_minimo:
            continue
        firma = (
            xml_de(p._p.find(qn("w:pPr"))),
            xml_de(p.runs[0]._r.find(qn("w:rPr"))),
        )
        cuenta[firma] += 1
        ejemplo.setdefault(firma, p)

    if not cuenta:
        return None

    return ejemplo[cuenta.most_common(1)[0][0]]


def aplicar_modelo(parrafo, modelo) -> None:
    if modelo is None:
        return

    p = parrafo._p
    viejo = p.find(qn("w:pPr"))
    nuevo = copy.deepcopy(modelo._p.find(qn("w:pPr")))

    if viejo is not None:
        p.replace(viejo, nuevo)
    else:
        p.insert(0, nuevo)

    modelo_rPr = modelo.runs[0]._r.find(qn("w:rPr")) if modelo.runs else None
    for run in parrafo.runs:
        actual = run._r.find(qn("w:rPr"))
        if actual is not None:
            run._r.remove(actual)
        if modelo_rPr is not None:
            run._r.insert(0, copy.deepcopy(modelo_rPr))


def reemplazar_texto(parrafo, viejo: str, nuevo: str) -> bool:
    completo = "".join(r.text for r in parrafo.runs)
    if viejo not in completo:
        return False
    parrafo.runs[0].text = completo.replace(viejo, nuevo)
    for r in parrafo.runs[1:]:
        r.text = ""
    return True


def main(ruta: Path, capturas: Path) -> int:
    doc = Document(str(ruta))

    modelos = {
        "Normal": modelo_de(doc, "Normal", 200),
        "Heading 1": modelo_de(doc, "Heading 1", 8),
        "Heading 2": modelo_de(doc, "Heading 2", 8),
    }

    for estilo, modelo in modelos.items():
        if modelo is None:
            fallos.append(f"Sin modelo de formato para «{estilo}»")
        else:
            pPr = modelo._p.find(qn("w:pPr"))
            jc = pPr.find(qn("w:jc")) if pPr is not None else None
            if jc is not None and jc.get(qn("w:val")) == "center":
                fallos.append(f"ABORTADO: el modelo de «{estilo}» está centrado (sería la portada)")

    if fallos:
        for f in fallos:
            print("  !! ", f)
        return 1

    # ---------------------------------------------------------- punto de corte
    ancla = None
    for p in doc.paragraphs:
        if p.style.name == "Heading 1" and p.text.strip().startswith("XX. CONCLUSIONES"):
            ancla = p
            break

    if ancla is None:
        print("  !!  No se encontró «XX. CONCLUSIONES»; no hay dónde insertar.")
        return 1

    # ------------------------------------------------------------- utilidades
    # El documento original llega a la Figura 53, y esas tres viven en los
    # anexos. Numerar desde 51 producia DUPLICADOS: dos «Figura 51», dos 52 y
    # dos 53. Lo estricto seria renumerar las de los anexos, porque los
    # capitulos nuevos van antes; se prefiere continuar desde 54 para no tocar
    # referencias cruzadas ya existentes. Una numeracion que no sigue el orden
    # de pagina es una imperfeccion; dos figuras con el mismo numero es un
    # error.
    # Lo mismo ocurre con las tablas: el documento original llega a la Tabla
    # 63, de modo que las nuevas empiezan en la 64.
    figuras = {"n": 54}

    def parrafo(texto: str, estilo: str = "Normal"):
        nuevo = ancla.insert_paragraph_before(texto)
        nuevo.style = doc.styles[estilo]
        aplicar_modelo(nuevo, modelos[estilo])
        return nuevo

    def figura(archivo: str, pie: str, ancho: float = 5.8) -> None:
        """Imagen centrada con su pie numerado y su autoría."""
        ruta_imagen = capturas / archivo
        if not ruta_imagen.exists():
            fallos.append(f"Falta la captura {archivo}")
            return

        contenedor = ancla.insert_paragraph_before()
        contenedor.alignment = WD_ALIGN_PARAGRAPH.CENTER
        contenedor.add_run().add_picture(str(ruta_imagen), width=Inches(ancho))

        n = figuras["n"]
        figuras["n"] += 1

        leyenda = ancla.insert_paragraph_before(
            f"Figura {n}. {pie} Fuente: Elaboración propia."
        )
        leyenda.style = doc.styles["Normal"]
        aplicar_modelo(leyenda, modelos["Normal"])
        leyenda.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in leyenda.runs:
            r.font.size = None
            r.italic = True

        cambios.append(f"Figura {n}: {archivo}")

    def tabla(encabezados: list[str], filas: list[list[str]], titulo: str, numero: int) -> None:
        """Tabla con el mismo aspecto que las del documento.

        Las 78 tablas existentes usan el estilo «Normal Table» y su cabecera
        verde viene de formato DIRECTO sobre las celdas, no de un estilo. Por
        eso no basta con asignar un estilo: hay que CLONAR una tabla real del
        documento con el mismo número de columnas y sustituirle el contenido.
        Construirla desde cero daría una rejilla gris que delataría a simple
        vista qué capítulos se añadieron después.
        """
        plantilla = next(
            (t for t in doc.tables if len(t.columns) == len(encabezados) and len(t.rows) >= 2),
            None,
        )

        if plantilla is None:
            fallos.append(f"Sin tabla modelo de {len(encabezados)} columnas para «{titulo[:30]}»")
            return

        parrafo(f"Tabla {numero}", "Normal")
        rotulo = parrafo(titulo, "Normal")
        for r in rotulo.runs:
            r.italic = True

        nueva_tbl = copy.deepcopy(plantilla._tbl)
        ancla._p.addprevious(nueva_tbl)

        from docx.table import Table

        t = Table(nueva_tbl, ancla._parent)

        modelo_fila = copy.deepcopy(t.rows[1]._tr)

        # Se dejan solo cabecera y una fila, y se clona esa fila lo necesario.
        for fila in list(t.rows)[2:]:
            nueva_tbl.remove(fila._tr)

        while len(t.rows) - 1 < len(filas):
            nueva_tbl.append(copy.deepcopy(modelo_fila))

        def escribir(celda, texto: str) -> None:
            """Escribe conservando el formato del primer run de la celda."""
            parrafos = celda.paragraphs
            primero = parrafos[0]

            for extra in parrafos[1:]:
                extra._p.getparent().remove(extra._p)

            if primero.runs:
                primero.runs[0].text = texto
                for r in primero.runs[1:]:
                    r._r.getparent().remove(r._r)
            else:
                primero.add_run(texto)

        for i, h in enumerate(encabezados):
            escribir(t.rows[0].cells[i], h)

        for f, valores in enumerate(filas, start=1):
            for i, valor in enumerate(valores):
                escribir(t.rows[f].cells[i], str(valor))

        parrafo("Nota. Elaboración propia.", "Normal")
        cambios.append(f"Tabla {numero}: {titulo[:40]}")

    # =======================================================================
    #  XX. CORRECCIONES DEL APF2
    # =======================================================================
    parrafo("XX. CORRECCIONES DEL APF2", "Heading 1")
    parrafo(
        "Este capítulo recoge las observaciones detectadas sobre el avance anterior y el "
        "estado en que queda cada una. Se incluyen también las que siguen abiertas: un "
        "cuadro que solo recogiera lo corregido no sería un cuadro de observaciones, sino "
        "una lista de logros."
    )

    parrafo("20.1. Observaciones subsanadas", "Heading 2")
    tabla(
        ["N.º", "Observación", "Corrección aplicada", "Evidencia"],
        [
            ["O-A1", "El intérprete PHP quedó expuesto a CVE-2026-109618 y CVE-2026-109619, ejecución remota de código sin autenticar (CVSS 9.2)", "Imagen base y entorno local actualizados a PHP 8.4.27; comprobador que aborta la construcción y el arranque sobre un intérprete vulnerable; endurecimiento en docker/seguridad.ini", "Commit cf98072"],
            ["O-A2", "El recuento de la suite aparecía con tres cifras distintas en el mismo documento", "Sincronizado con la ejecución real en el informe, el README y el documento", "Commit e995733"],
            ["O-A3", "El defecto D-05 reapareció: las descargas de comprobantes no comprobaban propiedad", "Regla trasladada a ComprobanteElectronicoPolicy, aplicada por igual a descargas y portal", "Commit 2315149 · 8 pruebas"],
            ["O-A4", "Figuraba un colaborador que no correspondía en la lista de contributors", "Verificado que no procedía del historial; recálculo forzado del índice", "Panel de contributors"],
            ["O-A5", "El modal «Editar producto» mostraba el icono de imagen rota", "Dos fallos: la ruta no se resolvía contra el origen de la API y no había respaldo. Ambos en ImagenConRespaldo.vue", "Commit 4fea5e5 · Figura 59"],
            ["O-A6", "Cerrar sesión no invalidaba la sesión hasta pasados cinco segundos de cuenta atrás decorativa", "La sesión se cierra primero; la pantalla informa del estado real", "Commit 454a3ac · Figura 58"],
            ["O-A7", "El frontend no disponía de herramienta de pruebas", "Vitest y Vue Test Utils configurados; 34 pruebas de componente", "Commit 3a653b7"],
            ["O-A8", "El apartado 18.3 declaraba 4 524 unidades vendibles frente a 4 501 de la Tabla 62", "Medido: 4 501 en las tres formas de contarlo, con cero unidades vencidas. La cifra 4 524 era incorrecta", "Consulta del 23/10/2026"],
            ["O-A9", "Las pantallas del Sprint 5 carecían de verificación visual", "Nueve capturas tomadas sobre el entorno en ejecución", "Figuras 54 a 62"],
        ],
        "Observaciones del APF2 y su subsanación",
        64,
    )

    parrafo("20.2. Observaciones que siguen abiertas", "Heading 2")
    tabla(
        ["N.º", "Observación", "Por qué sigue abierta"],
        [
            ["O-B1", "Las Figuras 43 y C.3 muestran ejecuciones de la suite con cifras ya superadas", "Son capturas: hay que volver a tomarlas. Reescribir el pie haría que mintiera sobre su propia imagen"],
            ["O-B2", "La Figura 11 rotula «26 migraciones versionadas» dentro del PNG; hoy son 46", "El dato está incrustado en la imagen; exige regenerar el diagrama"],
            ["O-B3", "Las cifras del portal público no coinciden entre Inicio, Sobre Nosotros y el pie", "Decisión de negocio sobre cuál es la correcta"],
            ["O-B4", "Los 884 productos del catálogo apuntan a un archivo de imagen que ya no existe", "Todas las imágenes responden 404. El respaldo lo resuelve visualmente, pero el catálogo sigue sin fotografías"],
        ],
        "Observaciones pendientes al cierre del APF3",
        65,
    )

    parrafo("20.3. Comparación antes y después", "Heading 2")
    parrafo(
        "Las capturas se tomaron sobre el entorno en ejecución, no sobre un prototipo, y "
        "cada una se acompañó de una comprobación sobre el propio documento HTML para "
        "verificar lo que no se aprecia a simple vista: dónde queda el foco y qué campos "
        "se marcan como inválidos."
    )

    figura("E11-login-escritorio.jpg",
           "Pantalla de acceso en resolución de escritorio, con el panel de marca que se "
           "desplaza al alternar entre acceso y registro.", 5.6)
    figura("E12-login-validaciones-y-foco.jpg",
           "Acceso con el formulario enviado en blanco: cada campo muestra su mensaje y el foco "
           "se sitúa en el primero que falla, sin que el botón se desplace.", 4.3)
    figura("E13-registro-medidor-contrasena.jpg",
           "Registro con el medidor de fortaleza de contraseña en funcionamiento.", 4.3)
    figura("E15-dialogo-confirmar-salida.jpg",
           "Diálogo de confirmación previo al cierre de sesión. El foco entra en la opción "
           "segura, no en la destructiva.", 4.3)
    figura("E16-logout-despedida-rol.jpg",
           "Cierre de sesión: la despedida reconoce el nombre y el rol de la cuenta, y la "
           "sesión ya está cerrada cuando aparece la cuenta atrás.", 4.3)
    figura("E17-modal-editar-producto-respaldo.jpg",
           "Modal «Editar producto» tras la corrección: cuando el archivo de imagen no existe "
           "se muestra un respaldo en lugar del icono de imagen rota del navegador.", 4.3)

    # =======================================================================
    #  XXI. GESTIÓN DE SOLICITUDES DE SERVICIO
    # =======================================================================
    parrafo("XXI. GESTIÓN DE SOLICITUDES DE SERVICIO", "Heading 1")
    parrafo(
        "Una solicitud de servicio es una petición habitual de un usuario para obtener "
        "información, un acceso o una acción estándar ya definida. No es lo mismo que un "
        "incidente, que es algo que debería funcionar y está fallando, ni que un cambio, "
        "que modifica el servicio. Confundirlos tiene coste real: una caída atendida como "
        "petición espera turno, y una petición tratada como incidente dispara una urgencia "
        "que no existe."
    )

    parrafo("21.1. Ciclo de atención", "Heading 2")
    tabla(
        ["Paso", "Qué ocurre", "Dónde queda registrado"],
        [
            ["1. Registrar", "La petición entra como issue con su código SOL-xxx. La issue es el ticket", "Issues #14 a #19"],
            ["2. Clasificar", "Se determina si es de información, de acceso o estándar", "Etiquetas del repositorio"],
            ["3. Validar", "Se comprueba que es una solicitud y que los datos están completos", "Campos del ticket"],
            ["4. Autorizar", "Si requiere aprobación, se solicita. Si la aprobación no corresponde, se rechaza o se escala; el paso no se omite", "Campo Aprobación del tablero"],
            ["5. Atender", "El responsable ejecuta la acción", "Campo Responsable de atención"],
            ["6. Verificar", "Se comprueba el resultado antes de darlo por hecho", "Campo Evidencia"],
            ["7. Cerrar", "Se cierra con su evidencia. Una solicitud sin evidencia no se cierra", "Estado final"],
        ],
        "Ciclo de atención de una solicitud de servicio",
        66,
    )

    parrafo("21.2. Catálogo de solicitudes atendidas", "Heading 2")
    tabla(
        ["Código", "Tipo", "Descripción", "Aprobación", "Estado final"],
        [
            ["SOL-001", "Información", "Cómo generar el reporte de ventas del mes", "No requiere", "Cerrada"],
            ["SOL-002", "Información", "El catálogo muestra menos unidades que el anaquel", "No requiere", "Cerrada"],
            ["SOL-003", "Acceso", "Alta de cuenta para vendedora nueva", "Propietaria aprueba", "Cerrada"],
            ["SOL-004", "Acceso", "Cliente solicita acceso al módulo de reportes", "Propietaria deniega", "RECHAZADA"],
            ["SOL-005", "Estándar", "Restablecer contraseña de vendedora", "Propietaria aprueba", "Cerrada"],
            ["SOL-006", "Estándar", "Respaldo de la base antes de una carga masiva", "Propietaria aprueba", "Cerrada"],
        ],
        "Catálogo de solicitudes de servicio del sistema",
        67,
    )

    parrafo(
        "SOL-004 es el caso que justifica el cuarto paso del ciclo. Un cliente pidió acceso "
        "a los reportes de ventas; conceder lo que se pide es lo cómodo, y el procedimiento "
        "existe precisamente para los casos en que no corresponde. Se denegó por el "
        "principio de mínimo privilegio, porque los reportes contienen márgenes, costos de "
        "compra y datos de otros clientes, y se ofreció la alternativa que sí le pertenece: "
        "su propio historial de pedidos y sus comprobantes. La evidencia no es una captura: "
        "es el propio sistema respondiendo 401 y 403 a un token de rol cliente, con pruebas "
        "automatizadas que fijan esa regla para que no se afloje por comodidad."
    )

    figura("E18-github-tablero-itsm.jpg",
           "Tablero de solicitudes de servicio en GitHub Projects con las seis solicitudes "
           "registradas y cerradas.", 5.8)
    figura("E19-github-sol004-campos.jpg",
           "Ficha de la solicitud SOL-004 con sus campos de seguimiento: tipo, prioridad, "
           "estado, aprobación y fecha.", 5.8)

    parrafo(
        "Conviene advertir una diferencia entre dos campos que parecen el mismo. El estado "
        "propio de GitHub pasa a «Done» en cuanto la issue se cierra, de modo que SOL-004 "
        "figura allí como completada aunque fue rechazada. El campo que manda en este "
        "procedimiento es «Estado solicitud», porque una solicitud denegada también se "
        "cierra, y cerrarla no significa haberla concedido."
    )

    # =======================================================================
    #  XXII. CONTROL DE CAMBIOS
    # =======================================================================
    parrafo("XXII. PROCEDIMIENTO DE CONTROL DE CAMBIOS", "Heading 1")
    parrafo(
        "El flujo es registrar, evaluar impacto y riesgo, autorizar, preparar el plan de "
        "vuelta atrás, ejecutar, verificar y cerrar documentando. El plan de vuelta atrás "
        "no es opcional: un cambio sin forma de deshacerlo no se autoriza. En este proyecto "
        "es barato, porque la plataforma conserva los despliegues anteriores y uno se "
        "reactiva desde el panel, de modo que no hay excusa para omitirlo."
    )

    parrafo("22.1. Clasificación de los cambios", "Heading 2")
    tabla(
        ["Tipo", "Cuándo se aplica", "Quién autoriza", "Ejemplo real del proyecto"],
        [
            ["Estándar", "Repetitivo y con resultado conocido", "Preaprobado", "Actualizar una dependencia menor"],
            ["Normal", "Afecta a funcionalidad o a datos", "Propietaria", "Portal del cliente del Sprint 5"],
            ["De emergencia", "Riesgo activo; esperar cuesta más que actuar", "Se ejecuta y se documenta a continuación", "Parche de PHP del 21/10/2026"],
        ],
        "Clasificación de cambios aplicada en el proyecto",
        68,
    )

    parrafo("22.2. Caso de estudio: cambio de emergencia CAM-001", "Heading 2")
    tabla(
        ["Campo", "Contenido"],
        [
            ["Código", "CAM-001"],
            ["Fecha", "21 de octubre de 2026"],
            ["Tipo", "De emergencia"],
            ["Origen", "Aviso de seguridad de php.net: CVE-2026-109618 y CVE-2026-109619, CVSS v4.0 9.2, ejecución remota de código sin autenticar"],
            ["Impacto si no se actúa", "El aviso nombra a FrankenPHP, que es el servidor de esta API. El fallo se dispara durante el análisis de la petición, antes de que se ejecute una sola línea de la aplicación, de modo que ni la autenticación ni los roles intervienen"],
            ["Alcance", "Imagen base del contenedor, configuración del intérprete en producción y entorno local de desarrollo"],
            ["Autorización", "Propietaria, con el riesgo explicado. Por tratarse de una emergencia se ejecutó primero y se documentó después"],
            ["Plan de vuelta atrás", "Reactivar el despliegue anterior desde el panel. Riesgo asumido: volvería a quedar expuesto, de modo que solo se usaría ante una caída del servicio"],
            ["Ejecución", "Se verificó que la etiqueta de la imagen ya apuntaba a la versión corregida comparando identificadores; se redesplegó con la caché limpia, porque de lo contrario la capa almacenada habría reutilizado la imagen anterior y el despliegue habría parecido correcto sin estarlo; se añadió el endurecimiento del intérprete"],
            ["Verificación", "El servicio responde correctamente, la versión del intérprete se leyó del propio servidor y el catálogo permanece íntegro en 884 productos"],
            ["Estado", "Cerrado"],
        ],
        "Registro del cambio de emergencia CAM-001",
        69,
    )

    parrafo(
        "Una decisión merece quedar escrita porque fue contraria a la recomendación del "
        "propio aviso. Entre las mitigaciones sugeridas figuraba suprimir el límite de "
        "memoria del intérprete. No se aplicó: esa recomendación presupone un límite del "
        "contenedor por debajo que detenga únicamente al proceso afectado, mientras que "
        "aquí el contenedor es el proceso. Suprimir el límite habría convertido un error "
        "acotado en la caída del servicio, y con el intérprete ya corregido el remedio era "
        "peor que la enfermedad."
    )

    parrafo(
        "El cambio dejó además una salvaguarda. La etiqueta de la imagen se mantiene "
        "flotante a propósito, para que la siguiente corrección llegue sola al reconstruir; "
        "lo que una etiqueta flotante no permite es demostrar qué versión se ejecuta, y eso "
        "lo cubre un comprobador que aborta la construcción y el arranque sobre un "
        "intérprete vulnerable. Usa una tabla de versiones mínimas por rama, porque "
        "comparar contra un único número daría un falso negativo con una rama posterior "
        "que todavía no estuviera corregida."
    )

    # =======================================================================
    #  XXIII. CAPACIDAD Y RENDIMIENTO
    # =======================================================================
    parrafo("XXIII. GESTIÓN DE CAPACIDAD Y RENDIMIENTO", "Heading 1")
    parrafo(
        "Todas las cifras de este capítulo se midieron el 23 de octubre de 2026. Se indican "
        "con su método de obtención para que cualquiera pueda repetir la medición y obtener "
        "la suya."
    )

    tabla(
        ["Indicador", "Método", "Umbral", "Resultado"],
        [
            ["Respuesta del catálogo en producción", "Siete peticiones consecutivas; se descarta la primera por el arranque en frío", "2 000 ms (RNF03)", "Mediana 382 ms; mínimo 359; máximo 434"],
            ["Salto de maquetación (CLS)", "PerformanceObserver sobre las vistas públicas", "0.1 (umbral de «bueno»)", "0.015 en catálogo; 0.013 a 0.015 en carrito y contacto"],
            ["Suite de pruebas del servidor", "Ejecución completa sobre PostgreSQL 16", "Totalidad en verde", "165 pruebas, 497 aserciones, 14.05 s"],
            ["Suite de componentes", "Ejecución completa con Vitest", "Totalidad en verde", "34 pruebas, 1.40 s"],
            ["Compilación del cliente", "Compilación de producción", "Sin avisos", "6.12 s"],
            ["Peso entregado al navegador", "Medición sobre los archivos compilados", "—", "Carga inicial de 189 kB, 71 kB comprimidos"],
        ],
        "Indicadores de capacidad y rendimiento medidos",
        70,
    )

    parrafo("23.1. Capacidad del sistema", "Heading 2")
    tabla(
        ["Dimensión", "Valor"],
        [
            ["Productos en catálogo", "884"],
            ["Unidades vendibles", "4 501"],
            ["Lotes registrados", "886"],
            ["Rutas de la API", "119"],
            ["Migraciones versionadas", "46"],
            ["Modelos, controladores y servicios", "23, 22 y 10"],
        ],
        "Capacidad actual del sistema",
        71,
    )

    parrafo("23.2. Límites conocidos", "Heading 2")
    parrafo(
        "El entorno desplegado funciona sobre un plan gratuito, con un único contenedor que "
        "se suspende por inactividad: la primera petición tras la suspensión tarda del orden "
        "de segundos, razón por la cual la medición descarta la primera muestra. Para una "
        "operación real haría falta un plan que no suspenda el servicio."
    )
    parrafo(
        "Conviene decir con claridad lo que no se ha medido: no se ha realizado una prueba "
        "de carga. Las cifras anteriores describen el tiempo de respuesta con un único "
        "usuario, no el comportamiento con varios simultáneos. Afirmar que el sistema "
        "soporta una cantidad determinada de usuarios sin haberlo comprobado sería inventarlo."
    )

    # =======================================================================
    #  XXIV. DESPLIEGUE DE LA VERSIÓN FINAL
    # =======================================================================
    parrafo("XXIV. DESPLIEGUE DE LA VERSIÓN FINAL", "Heading 1")
    tabla(
        ["Componente", "Plataforma", "Estado verificado el 23/10/2026"],
        [
            ["Interfaz de usuario", "Vercel", "En servicio, responde correctamente"],
            ["API", "Render, contenedor con PHP 8.4.27", "En servicio; despliegue automático desde la rama principal"],
            ["Base de datos", "Neon, PostgreSQL 16", "884 productos y 46 migraciones aplicadas"],
            ["Control de acceso", "Comprobado sin credenciales", "Las rutas del portal del cliente responden 401 sin sesión"],
        ],
        "Estado del despliegue de la versión final",
        72,
    )

    figura("E14-portal-cliente-escritorio.jpg",
           "Portal del cliente con datos reales: cifras agregadas en la base de datos, "
           "historial de pedidos con su estado y acceso a los comprobantes.", 5.8)

    # =======================================================================
    #  XXV. RETROSPECTIVA DEL SPRINT 6
    # =======================================================================
    parrafo("XXV. RETROSPECTIVA DEL SPRINT 6", "Heading 1")

    parrafo("25.1. Qué mantener", "Heading 2")
    tabla(
        ["Práctica", "Qué evitó"],
        [
            ["Leer el código antes de afirmar qué le falta", "Dos de los cinco defectos atribuidos a la pantalla de acceso ya estaban resueltos; una búsqueda por palabra clave en un solo archivo no los encontró"],
            ["Verificar leyendo el resultado, no el eco del propio comando", "Los campos del tablero se dieron por guardados cuando estaban vacíos, porque el guion mostraba su intención y silenciaba los errores"],
            ["Trasladar la regla a un único sitio en lugar de repetirla", "El defecto D-05 reapareció precisamente porque la primera vez se corrigió dentro de un método concreto"],
            ["Reservar el espacio antes de que lleguen los datos", "El salto de maquetación permanece en 0.015 con tarjetas, historial y comprobantes nuevos"],
        ],
        "Prácticas que conviene mantener",
        73,
    )

    parrafo("25.2. Qué dejar de hacer", "Heading 2")
    tabla(
        ["Práctica", "Qué costó"],
        [
            ["Diagnosticar por búsqueda de texto lo que exige leer", "Dos informes de alcance equivocados en sesiones consecutivas"],
            ["Silenciar la salida de un comando que puede fallar", "Treinta ediciones del tablero fallaron sin que nadie se enterara"],
            ["Dar por protegido un punto de acceso porque su grupo exige sesión", "Exigir identidad no es exigir propiedad: las descargas de comprobantes quedaron abiertas a cualquier cuenta"],
            ["Declarar cifras que no se acaban de medir", "El recuento de la suite llegó a aparecer con tres valores distintos en el mismo documento"],
        ],
        "Prácticas que causaron retrabajo",
        74,
    )

    parrafo("25.3. Qué empezar a hacer", "Heading 2")
    parrafo(
        "Capturar la pantalla y no solo compilarla. El recorrido visual de este avance no "
        "sirvió únicamente para documentar lo hecho: detectó una regresión introducida en "
        "el propio Sprint. Al reescribir el portal del cliente desapareció el botón de "
        "cerrar sesión que tenía la versión anterior, de modo que se entraba al panel y no "
        "había forma de salir salvo escribiendo la dirección a mano. Ninguna de las treinta "
        "y cuatro pruebas de componente podía detectarlo, porque todas comprueban que lo "
        "que existe se comporta correctamente y ninguna comprueba que algo que debería "
        "existir no falte."
    )
    parrafo(
        "Realizar una prueba de carga antes de hablar de capacidad; revisar los puntos de "
        "acceso ya existentes con el mismo criterio que los nuevos, porque el defecto "
        "reaparecido vivía en una ruta que nadie había vuelto a mirar; y cerrar las "
        "observaciones que permanecen abiertas, empezando por las figuras que muestran "
        "cifras ya superadas."
    )

    parrafo("25.4. Compromisos para el siguiente Sprint", "Heading 2")
    tabla(
        ["N.º", "Acción", "Indicador de éxito"],
        [
            ["A-08", "Recapturar las figuras que muestran cifras superadas", "Los pies coinciden con lo que muestra cada imagen"],
            ["A-09", "Regenerar el diagrama del modelo de datos con el número correcto de migraciones", "El rótulo del diagrama coincide con el repositorio"],
            ["A-10", "Cargar las imágenes reales del catálogo", "Los productos muestran su fotografía y no el respaldo"],
            ["A-11", "Auditar las rutas de la API sin prueba de autorización", "Lista revisada, con o sin hallazgos"],
            ["A-12", "Ejecutar una prueba de carga sobre el entorno desplegado", "Cifra de usuarios simultáneos medida y documentada"],
        ],
        "Acciones comprometidas para el Sprint 7",
        75,
    )

    # ------------------------------------------------- renumerar el cierre
    for p in doc.paragraphs:
        if p.style.name != "Heading 1":
            continue
        if reemplazar_texto(p, "XX. CONCLUSIONES", "XXVI. CONCLUSIONES"):
            cambios.append("Renumerado: XX. CONCLUSIONES → XXVI")
        elif reemplazar_texto(p, "XXI. REFERENCIAS", "XXVII. REFERENCIAS"):
            cambios.append("Renumerado: XXI. REFERENCIAS → XXVII")

    doc.save(str(ruta))

    print("CAMBIOS APLICADOS")
    for c in cambios:
        print("  OK  ", c)
    if fallos:
        print("\nNO APLICADOS")
        for f in fallos:
            print("  !!  ", f)
    print(f"\nTotal: {len(cambios)} cambios, {len(fallos)} sin aplicar.")
    print("Pulsa F9 en Word para actualizar el índice automático.")
    return 1 if fallos else 0


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print(__doc__)
        sys.exit(2)
    sys.exit(main(Path(sys.argv[1]), Path(sys.argv[2])))
